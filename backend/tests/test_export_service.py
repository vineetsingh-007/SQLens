import io
import pytest
from fastapi.testclient import TestClient
from docx import Document
import openpyxl

from app.main import app
from app.schemas.export import ExportRequest
from app.services.export_service import generate_excel, generate_docx, generate_pdf

client = TestClient(app)


def test_generate_excel_success():
    req = ExportRequest(
        columns=["product", "revenue", "in_stock"],
        rows=[
            {"product": "Laptop", "revenue": 125000, "in_stock": True},
            {"product": "Phone", "revenue": 95000, "in_stock": False},
            {"product": "Tablet", "revenue": None, "in_stock": None},
        ],
        question="What are product revenues?",
        insight="Laptop has the highest revenue.",
        truncated=False
    )
    res_bytes = generate_excel(req)
    assert isinstance(res_bytes, bytes)
    assert len(res_bytes) > 0

    # Verify openpyxl can load the generated Excel stream
    wb = openpyxl.load_workbook(io.BytesIO(res_bytes))
    assert "Results" in wb.sheetnames
    ws = wb["Results"]
    assert ws.cell(row=1, column=1).value == "product"
    assert ws.cell(row=2, column=1).value == "Laptop"


def test_generate_docx_success():
    req = ExportRequest(
        columns=["city", "customer_count"],
        rows=[
            {"city": "Pune", "customer_count": 42},
            {"city": "Mumbai", "customer_count": 128},
        ],
        question="Customer count by city",
        insight="Mumbai leads in total customer base.",
        truncated=True
    )
    res_bytes = generate_docx(req)
    assert isinstance(res_bytes, bytes)
    assert len(res_bytes) > 0

    # Verify docx can load the generated Word stream
    doc = Document(io.BytesIO(res_bytes))
    full_text = " ".join([p.text for p in doc.paragraphs])
    assert "SQLens Query Result" in full_text
    assert "Customer count by city" in full_text
    assert "Note: Export contains the 2 rows returned by SQLens." in full_text


def test_generate_pdf_success():
    req = ExportRequest(
        columns=["id", "name", "email", "category", "price", "status"],
        rows=[
            {"id": 1, "name": "Item A", "email": "a@ex.com", "category": "Gadgets", "price": 99.99, "status": "Active"},
            {"id": 2, "name": "Item B", "email": "b@ex.com", "category": "Widgets", "price": 49.50, "status": "Inactive"},
        ],
        question="Show catalog products",
        insight="Catalog contains active gadgets.",
        truncated=False
    )
    res_bytes = generate_pdf(req)
    assert isinstance(res_bytes, bytes)
    assert res_bytes.startswith(b"%PDF-")


def test_export_endpoints():
    payload = {
        "columns": ["dept", "budget"],
        "rows": [{"dept": "Engineering", "budget": 500000}],
        "question": "Engineering budget",
        "insight": "Approved for Q3",
        "truncated": False
    }

    # Test Excel
    res_excel = client.post("/api/export/excel", json=payload)
    assert res_excel.status_code == 200
    assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # Test Word
    res_docx = client.post("/api/export/docx", json=payload)
    assert res_docx.status_code == 200
    assert res_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    # Test PDF
    res_pdf = client.post("/api/export/pdf", json=payload)
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"


def test_export_empty_result_bad_request():
    payload = {
        "columns": [],
        "rows": [],
        "question": "Empty query"
    }

    res = client.post("/api/export/excel", json=payload)
    assert res.status_code == 400
    assert "Nothing to export" in res.json()["detail"]


def test_export_visualization_success():
    mock_png_b64 = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    payload = {
        "columns": ["category", "sales"],
        "rows": [{"category": "Tech", "sales": 500}],
        "question": "Sales by category",
        "insight": "Tech category lead sales",
        "include_visualization": True,
        "chart_image_base64": mock_png_b64
    }

    # Test PDF with visualization
    res_pdf = client.post("/api/export/pdf", json=payload)
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert res_pdf.content.startswith(b"%PDF-")

    # Test Word with visualization
    res_docx = client.post("/api/export/docx", json=payload)
    assert res_docx.status_code == 200
    assert res_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

